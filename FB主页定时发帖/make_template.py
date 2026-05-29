"""生成 posts.xlsx 模板。一次性脚本，跑完就能删。"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

wb = Workbook()
ws = wb.active
ws.title = "Sheet1"

headers = ["文案", "定时时间", "素材1", "素材2", "素材3", "素材4", "素材5",
           "素材6", "素材7", "素材8", "素材9", "素材10"]
ws.append(headers)

for c in ws[1]:
    c.font = Font(bold=True)
    c.fill = PatternFill("solid", fgColor="DDEBF7")
    c.alignment = Alignment(horizontal="center")

ws.append([
    "今日上新，点击主页了解详情~",
    "2026-06-05 10:00",
    "demo_image_1.jpg", "", "", "", "", "", "", "", "", "",
])
ws.append([
    "多图测试，第二条是 3 图轮播",
    "2026-06-06 14:00",
    "a.jpg", "b.jpg", "c.jpg", "", "", "", "", "", "", "",
])
ws.append([
    "视频帖示例（视频帖只能放 1 个）",
    "",
    "intro.mp4", "", "", "", "", "", "", "", "", "",
])

ws.column_dimensions["A"].width = 50
ws.column_dimensions["B"].width = 20
for col in "CDEFGHIJKL":
    ws.column_dimensions[col].width = 22

wb.save("posts_template.xlsx")
print("已生成 posts_template.xlsx")
